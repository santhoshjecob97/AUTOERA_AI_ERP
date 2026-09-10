import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'f:\autoeraaisaas-main\AUTOERA PITCH DOCUMENTS\AutoEra_AI_Product_Architecture_2026.docx')

print("=== PARAGRAPHS OVERVIEW ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    style_name = getattr(p.style, 'name', '') if p.style else ''
    if txt and (len(txt) < 120 or 'Heading' in style_name or 'Title' in style_name):
        print(f"P{i} [{style_name}]: {txt}")

print(f"\n=== TABLES OVERVIEW ({len(doc.tables)} tables) ===")
for i, t in enumerate(doc.tables):
    if len(t.rows) > 0 and len(t.rows[0].cells) > 0:
        row0 = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
        headers_str = " | ".join(row0[:5])
        print(f"Table {i+1} ({len(t.rows)} rows): {headers_str[:120]}")
